"""固定页解锁首帧、重排恢复与原生导出共用几何。"""

import copy
from io import BytesIO

import pytest
from pptx import Presentation
from pydantic import ValidationError

from app.domain.flex_edit import build_flex_tree_from_fixed
from app.domain.flex_layout import FlexContainer, FlexLeaf, clear_rect_overrides, move_leaf_preserve_size
from app.domain.geometry import Rect
from app.domain.layout import get_layout
from app.regression.corpus import load_corpus_decks
from app.domain.slide_geometry import placed_by_block_id
from app.llm.edit_tools import EditSession, build_edit_tools
from app.render.pptx import render_deck_to_pptx


def unlocked(slide):
    tree = build_flex_tree_from_fixed(get_layout(slide.layout_id), [b.model_dump(mode="json") for b in slide.blocks])
    return slide.model_copy(update={"layout_mode": "flex", "layout_tree": tree})


def leaves(tree):
    for child in tree.children:
        if isinstance(child, FlexLeaf):
            yield child
        else:
            yield from leaves(child)


@pytest.mark.parametrize("index", range(11))
def test_every_fixed_layout_unlock_preserves_exact_rect_and_content(index):
    deck = load_corpus_decks()["short"]
    assert len({s.layout_id for s in deck.slides}) == 11
    fixed = deck.slides[index]
    original = copy.deepcopy(fixed.model_dump())
    flex = unlocked(fixed)
    assert fixed.model_dump() == original
    assert flex.blocks == fixed.blocks
    actual, expected = placed_by_block_id(flex), placed_by_block_id(fixed)
    assert set(actual) == set(expected) == {b.id for b in fixed.blocks}
    for block_id, geometry in expected.items():
        assert actual[block_id].rect.model_dump() == geometry.rect.model_dump()
        assert actual[block_id].text_style == geometry.text_style


@pytest.mark.parametrize("rect", [
    {"x": float("nan"), "y": 0, "w": .1, "h": .1},
    {"x": 0, "y": 0, "w": float("inf"), "h": .1},
    {"x": .9, "y": 0, "w": .2, "h": .1},
    {"x": 0, "y": .9, "w": .1, "h": .2},
    {"x": 0, "y": 0, "w": 0, "h": .1},
])
def test_override_rejects_nonfinite_empty_or_off_canvas_rect(rect):
    with pytest.raises(ValidationError):
        FlexLeaf(id="leaf", block_id="block", rect_override=rect)


def test_clearing_and_manual_move_restore_tree_without_mutating_original():
    tree = FlexContainer(type="column", id="root", children=[
        FlexLeaf(id="a", block_id="a", rect_override=Rect(x=.1, y=.1, w=.5, h=.1)),
        FlexLeaf(id="b", block_id="b", rect_override=Rect(x=.1, y=.6, w=.5, h=.1)),
    ])
    original = tree.model_dump()
    assert all(leaf.rect_override is None for leaf in leaves(clear_rect_overrides(tree)))
    moved = move_leaf_preserve_size(tree, "b", "root", 0)
    assert moved is not None and [leaf.block_id for leaf in leaves(moved)] == ["b", "a"]
    assert all(leaf.rect_override is None for leaf in leaves(moved))
    assert tree.model_dump() == original


@pytest.mark.asyncio
@pytest.mark.parametrize("tool_name,args", [
    ("add_block", {"after_block_id": "b1", "block_type": "text", "text": "新增内容"}),
    ("delete_block", {"block_id": "b1"}),
    ("change_type", {"block_id": "b1", "new_type": "text", "text": "改成文字"}),
])
async def test_successful_structure_tool_clears_override(tool_name, args):
    slide = unlocked(load_corpus_decks()["short"].slides[1])
    assert any(leaf.rect_override is not None for leaf in leaves(slide.layout_tree))
    session = EditSession(blocks=list(slide.blocks), tree=slide.layout_tree, layout_mode="flex")
    tools = {tool.name: tool for tool in build_edit_tools(session)}
    args = dict(args)
    args["after_block_id" if tool_name == "add_block" else "block_id"] = slide.blocks[-1].id
    result = await tools[tool_name].ainvoke(args)
    assert result.startswith("已"), result
    assert all(leaf.rect_override is None for leaf in leaves(session.tree))


@pytest.mark.asyncio
async def test_text_replacement_and_rejected_structure_keep_first_frame():
    slide = unlocked(load_corpus_decks()["short"].slides[1])
    session = EditSession(blocks=list(slide.blocks), tree=slide.layout_tree, layout_mode="flex")
    tools = {tool.name: tool for tool in build_edit_tools(session)}
    original = session.tree.model_dump()
    assert (await tools["replace_text"].ainvoke({"block_id": slide.blocks[0].id, "text": "新标题"})).startswith("已")
    assert session.tree.model_dump() == original
    assert (await tools["delete_block"].ainvoke({"block_id": "missing-block"})).startswith("错误")
    assert session.tree.model_dump() == original


def _native_content(presentation):
    result = []
    for slide in presentation.slides:
        content = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text:
                key = ("text", shape.text)
            elif shape.has_table:
                key = ("table", tuple(cell.text for row in shape.table.rows for cell in row.cells))
            elif shape.has_chart:
                key = ("chart", tuple(series.name for series in shape.chart.series))
            else:
                continue
            content.append((key, (shape.left, shape.top, shape.width, shape.height)))
        result.append(content)
    return result


def test_unlock_native_pptx_text_table_chart_positions_match_fixed():
    fixed = load_corpus_decks()["short"]
    flex = fixed.model_copy(update={"slides": [unlocked(slide) for slide in fixed.slides]})
    fixed_pptx = Presentation(BytesIO(render_deck_to_pptx(fixed).getvalue()))
    flex_pptx = Presentation(BytesIO(render_deck_to_pptx(flex).getvalue()))
    assert _native_content(fixed_pptx) == _native_content(flex_pptx)
