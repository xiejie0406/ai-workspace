import uuid

from collections.abc import AsyncGenerator

import pytest

from pptx import Presentation

from pptx.enum.text import PP_ALIGN

from pptx.util import Pt

from pydantic import ValidationError

from app.domain.block_style import BlockStyle, merge_text_style, resolve_box

from app.domain.content import Deck, Slide, TextBlock

from app.domain.theme import get_theme

from app.domain.validation import validate_slide

from app.render.pptx import render_deck_to_pptx

def test_merge_text_style_passthrough() -> None:
    theme = get_theme("ivory")
    merged = merge_text_style(theme, "title", None)
    base = theme.text_style("title")
    assert merged.size_pt == base.size_pt
    assert merged.color == base.color
    assert merged.italic is False

def test_merge_text_style_overrides() -> None:
    theme = get_theme("ivory")
    style = BlockStyle(
        size_pt=28,
        color="accent",
        weight=700,
        italic=True,
    )
    merged = merge_text_style(theme, "title", style)
    assert merged.size_pt == 28
    assert merged.color == "accent"
    assert merged.weight == 700
    assert merged.italic is True
    assert theme.color(merged.color) == theme.palette.accent

def test_merge_text_style_hex_color() -> None:
    theme = get_theme("ivory")
    style = BlockStyle(color="#0D9488")
    merged = merge_text_style(theme, "body", style)
    assert merged.color == "#0D9488"
    assert theme.color(merged.color) == "#0D9488"

def test_block_style_rejects_out_of_range() -> None:
    with pytest.raises(ValidationError):
        BlockStyle(size_pt=100)
    with pytest.raises(ValidationError):
        BlockStyle(color="not-a-color")
    with pytest.raises(ValidationError):
        BlockStyle(weight=100)

def test_resolve_box_fill_and_border() -> None:
    theme = get_theme("midnight")
    box = resolve_box(
        theme,
        BlockStyle(
            fill="accent_soft",
            radius_pt=8,
            border_color="accent",
            border_width_pt=1.5,
            padding_pt=12,
        ),
    )
    assert box.fill == theme.palette.accent_soft
    assert box.radius_pt == 8
    assert box.has_border
    assert box.border_color == theme.palette.accent
    assert box.padding_pt == 12

def test_overflow_uses_merged_size() -> None:
    from app.domain.content import BulletsBlock

    theme = get_theme("ivory")
    slide = Slide(
        id="s1",
        layout_id="bullets",
        blocks=[
            TextBlock(
                id="t1",
                slot_id="title",
                text="这是一段很长很长很长很长很长很长很长很长的标题内容",
                style=BlockStyle(size_pt=72),
            ),
            BulletsBlock(id="b1", slot_id="body", items=["要点"]),
        ],
    )
    issues = validate_slide(slide, theme=theme)
    assert any(issue.severity == "warning" and "溢出" in issue.message for issue in issues)

def test_pptx_applies_block_style() -> None:
    theme = get_theme("ivory")
    deck = Deck(
        id="d1",
        title="样式导出",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="cover",
                blocks=[
                    TextBlock(
                        id="t1",
                        slot_id="title",
                        text="加粗居中",
                        style=BlockStyle(
                            size_pt=40,
                            weight=700,
                            italic=True,
                            align="center",
                            color="#112233",
                        ),
                    ),
                    TextBlock(
                        id="s1",
                        slot_id="subtitle",
                        text="副标题",
                    ),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck, theme=theme)
    presentation = Presentation(buffer)
    slide = presentation.slides[0]

    # 找到含「加粗居中」的段落
    hit = None
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs)
            if "加粗居中" in text:
                hit = paragraph
                break
        if hit is not None:
            break

    assert hit is not None
    assert hit.alignment == PP_ALIGN.CENTER
    run = hit.runs[0]
    assert run.font.size == Pt(40)
    assert run.font.bold is True
    assert run.font.italic is True
    from app.render.color import to_rgb

    assert run.font.color.rgb == to_rgb("#112233")
