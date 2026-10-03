"""实际坏 PPTX 经过原 verify、handler、runner 和机器 HTTP 失败回执。"""
import ast
from io import BytesIO
import json
from pathlib import Path
import httpx
import pytest
from pptx import Presentation

from app.core.config import Settings
from app.domain.sample import load_sample_deck
from app.render.verify import VerifyIssue, VerifyReport, verify_pptx
from app.worker.catalog import catalog_version
from app.worker.handlers import TaskHandlers
from app.worker.protocol import CenterClient, ClaimedTask
from app.worker.runner import Worker


async def execute_corrupted_notes(monkeypatch):
    from app.render import pptx as renderer
    deck = load_sample_deck().model_copy(deep=True)
    deck.slides = deck.slides[:1]
    deck.slides[0].speaker_notes = "原快照备注必须精确保留\nnotes-content-fixture"
    snapshot = {"project": {"id": deck.id, "title": deck.title, "theme_id": deck.theme_id},
        "slides": [slide.model_dump(mode="json") for slide in deck.slides]}
    task = ClaimedTask(taskId="fixture-export", attempt=1, leaseToken="fixture-lease", fencingToken=3,
        type="export", catalogVersion=catalog_version(), payload={"font_manifest_hash": "unavailable", "snapshot": snapshot})
    original_render = renderer.render_deck_to_pptx
    rendered = []
    def corrupt_after_real_render(*args, **kwargs):
        good = original_render(*args, **kwargs).getvalue()
        presentation = Presentation(BytesIO(good))
        presentation.slides[0].notes_slide.notes_text_frame.text = "被替换的备注"
        out = BytesIO(); presentation.save(out)
        rendered.append(out.getvalue())
        return BytesIO(rendered[-1])
    monkeypatch.setattr(renderer, "render_deck_to_pptx", corrupt_after_real_render)
    requests = []
    def respond(request):
        requests.append((request.url.path, json.loads(request.content)))
        return httpx.Response(200, json={"code": 200, "data": {"accepted": True, "leaseValid": True}})
    cfg = Settings(_env_file=None, center_url="http://fixture.invalid", worker_key="fixture-key")
    center = CenterClient(cfg, transport=httpx.MockTransport(respond))
    try:
        await Worker(cfg, center=center, handlers=TaskHandlers(center)).execute(task)
    finally:
        await center.close()
    assert len(rendered) == 1
    report = verify_pptx(rendered[0], deck)
    assert report.passed is False and [issue.check for issue in report.issues] == ["speaker_notes"]
    assert report.issues[0].slide_index == 1
    return requests, report, rendered[0]


@pytest.mark.asyncio
async def test_actual_bad_pptx_reports_native_verify_diagnostics_without_upload(monkeypatch, caplog):
    requests, report, _ = await execute_corrupted_notes(monkeypatch)
    assert [path.rsplit("/", 1)[-1] for path, _ in requests] == ["progress", "fail"]
    body = requests[-1][1]
    assert body["code"] == "PPTX_VERIFY_FAILED" and body["retryable"] is False
    assert body["verification_failure"] == {**report.model_dump(mode="json"),
        "diagnostics_truncated": False, "total_issues": 1}
    assert {key: body[key] for key in ("attempt", "leaseToken", "fencingToken")} == {
        "attempt": 1, "leaseToken": "fixture-lease", "fencingToken": 3}
    assert body["execution"]["usage_status"] == "not_applicable"
    assert "ExportVerificationFailed" in caplog.text and "notes-content-fixture" not in caplog.text
    assert "verification_failure" not in caplog.text and "原快照备注" not in caplog.text


def test_failure_report_uses_all_actual_verifier_check_codes():
    from app.worker.protocol import ExportVerificationFailed
    import app.render.verify as module
    tree = ast.parse(Path(module.__file__).read_text(encoding="utf-8"))
    actual = {kw.value.value for node in ast.walk(tree) if isinstance(node, ast.Call)
        for kw in node.keywords if kw.arg == "check" and isinstance(kw.value, ast.Constant)}
    assert ExportVerificationFailed.CHECKS == actual
    report = VerifyReport(passed=False, slide_count=40, expected_slide_count=40,
        issues=[VerifyIssue(check=code, slide_index=1, shape=None, message="实际诊断类型") for code in sorted(actual)])
    payload = ExportVerificationFailed(report).verification_failure
    assert {issue["check"] for issue in payload["issues"]} == actual
    assert payload["total_issues"] == len(actual) and payload["diagnostics_truncated"] is False


def test_diagnostic_limits_preserve_original_count_and_explicit_truncation():
    from app.worker.protocol import ExportVerificationFailed
    report = VerifyReport(passed=False, slide_count=40, expected_slide_count=40,
        issues=[VerifyIssue(check="content_integrity", slide_index=1, shape="🙂" * 201, message="🙂" * 1001) for _ in range(2001)])
    payload = ExportVerificationFailed(report).verification_failure
    assert payload["total_issues"] == 2001 and len(payload["issues"]) == 2000
    assert payload["diagnostics_truncated"] is True
    assert len(payload["issues"][0]["message"]) == 1000 and len(payload["issues"][0]["shape"]) == 200
    assert payload["issues"][0]["message"].endswith("…") and payload["issues"][0]["shape"].endswith("…")
    assert len(report.issues) == 2001 and len(report.issues[0].message) == 1001
    assert "🙂" not in str(ExportVerificationFailed(report))


def test_exact_diagnostic_limits_do_not_claim_truncation():
    from app.worker.protocol import ExportVerificationFailed
    report = VerifyReport(passed=False, expected_slide_count=1, issues=[
        VerifyIssue(check="open", slide_index=None, shape="🙂" * 200, message="🙂" * 1000)])
    payload = ExportVerificationFailed(report).verification_failure
    assert payload["diagnostics_truncated"] is False and payload["total_issues"] == 1
    assert payload["issues"][0] == report.issues[0].model_dump(mode="json")


@pytest.mark.parametrize("report", [VerifyReport(passed=True), VerifyReport(passed=False),
    VerifyReport(passed=False, issues=[VerifyIssue(check="supplier_body", message="private provider reply")]),
    {"passed": False, "issues": [{"message": "arbitrary exception"}]}])
def test_only_actual_failed_verify_report_can_enter_export_diagnostics(report):
    from app.worker.protocol import ExportVerificationFailed
    with pytest.raises(ValueError, match="内部诊断契约"):
        ExportVerificationFailed(report)


@pytest.mark.asyncio
async def test_existing_plain_fail_caller_does_not_emit_optional_diagnostics():
    requests = []
    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(200, json={"data": {"accepted": True}})
    center = CenterClient(Settings(_env_file=None, worker_key="fixture-key"), transport=httpx.MockTransport(respond))
    task = ClaimedTask(taskId="fixture", attempt=1, leaseToken="lease", fencingToken=1,
        type="parse_source", catalogVersion=catalog_version(), payload={})
    try:
        await center.fail(task, "INVALID_CONTENT", "固定公开错误", False)
        assert "verification_failure" not in requests[0] and "execution" not in requests[0]
    finally:
        await center.close()
