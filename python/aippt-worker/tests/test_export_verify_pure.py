"""可编辑性回读验证与项目导出接口。"""

from __future__ import annotations

import uuid

from collections.abc import AsyncGenerator

from io import BytesIO

import pytest

from pptx import Presentation

from pptx.util import Emu, Inches, Pt

from app.domain.content import BulletsBlock, Deck, Slide, TextBlock

from app.domain.geometry import CANVAS_HEIGHT_PT, CANVAS_WIDTH_PT

from app.domain.sample import load_sample_deck

from app.render.pptx import render_deck_to_pptx

from app.render.verify import verify_pptx

PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

def test_verify_sample_deck_passes() -> None:
    deck = load_sample_deck()
    payload = render_deck_to_pptx(deck).getvalue()
    report = verify_pptx(payload, deck)

    assert report.passed is True, [issue.message for issue in report.issues]
    assert report.slide_count == len(deck.slides)
    assert report.expected_slide_count == len(deck.slides)
    assert report.issues == []


@pytest.mark.parametrize("notes", [None, "", "备注正文\n第二段  保留空格\t缩进\n末段"])
def test_verify_speaker_notes_round_trip_from_actual_pptx(notes):
    deck=load_sample_deck().model_copy(deep=True);deck.slides=deck.slides[:1];deck.slides[0].speaker_notes=notes
    data=render_deck_to_pptx(deck).getvalue()
    report=verify_pptx(data,deck)
    assert report.passed, [issue.message for issue in report.issues]


@pytest.mark.parametrize("changed", ["", "备注被改字"])
def test_verify_rejects_deleted_or_changed_speaker_notes(changed):
    deck=load_sample_deck().model_copy(deep=True);deck.slides=deck.slides[:1];deck.slides[0].speaker_notes="原备注\n第二段"
    presentation=Presentation(render_deck_to_pptx(deck))
    presentation.slides[0].notes_slide.notes_text_frame.text=changed
    damaged=BytesIO();presentation.save(damaged)
    report=verify_pptx(damaged.getvalue(),deck)
    assert not report.passed and any(issue.check=="speaker_notes" and issue.slide_index==1 for issue in report.issues)

def test_verify_detects_page_count_mismatch() -> None:
    deck = load_sample_deck()
    short = deck.model_copy(update={"slides": deck.slides[:1]})
    payload = render_deck_to_pptx(short).getvalue()

    report = verify_pptx(payload, deck)
    assert report.passed is False
    assert any(issue.check == "page_count" for issue in report.issues)

def test_verify_detects_out_of_bounds_shape() -> None:
    deck = Deck(
        id="d1",
        title="越界",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["要点"]),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck)
    presentation = Presentation(buffer)
    slide = presentation.slides[0]
    # 故意放到画布外
    slide.shapes.add_textbox(
        Emu(Pt(CANVAS_WIDTH_PT + 40)),
        Emu(Pt(10)),
        Emu(Pt(100)),
        Emu(Pt(40)),
    )
    bad = BytesIO()
    presentation.save(bad)

    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check == "bounds" and issue.slide_index == 1 for issue in report.issues)

def test_verify_detects_missing_text() -> None:
    deck = Deck(
        id="d1",
        title="缺字",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="真实标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["真实要点"]),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck)
    presentation = Presentation(buffer)
    # 清空文本框，制造「文字丢失」
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            shape.text_frame.clear()
    bad = BytesIO()
    presentation.save(bad)

    # 源 Deck 仍要求这些文字存在
    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check in {"text_frame", "content_integrity"} for issue in report.issues)

def test_verify_detects_full_page_picture() -> None:
    deck = Deck(
        id="d1",
        title="整页图",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["要点"]),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck)
    presentation = Presentation(buffer)
    # 1x1 PNG
    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00"
        b"\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    presentation.slides[0].shapes.add_picture(
        BytesIO(png),
        Emu(0),
        Emu(0),
        Emu(Pt(CANVAS_WIDTH_PT)),
        Emu(Pt(CANVAS_HEIGHT_PT)),
    )
    bad = BytesIO()
    presentation.save(bad)

    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check == "full_page_picture" for issue in report.issues)

def test_verify_detects_wrong_canvas_size() -> None:
    deck = Deck(
        id="d1",
        title="尺寸",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["要点"]),
                ],
            )
        ],
    )
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    bad = BytesIO()
    presentation.save(bad)

    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check == "canvas_size" for issue in report.issues)
