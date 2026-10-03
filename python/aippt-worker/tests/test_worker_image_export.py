import hashlib
from io import BytesIO
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.domain.content import Deck, Slide, TextBlock, BulletsBlock, ImageBlock
from app.render.verify import verify_pptx
from app.worker.catalog import catalog_version
from app.worker.handlers import TaskHandlers
from app.worker.protocol import ClaimedTask
from app.images.validate import ImageRejected


class Center:
    def __init__(self, data): self.data = data; self.uploaded = None
    async def read_asset(self, task, asset_id, limit): return self.data
    async def upload_asset(self, task, data, filename, mime, role):
        self.uploaded = data
        return {"asset_id": "export-fixture"}


def deck_and_task():
    slide = Slide(id="slide", layout_id="image-left", blocks=[
        TextBlock(id="title", slot_id="title", text="实际配图"),
        BulletsBlock(id="body", slot_id="body", items=["原始图片保留", "原生嵌入与裁剪"]),
        ImageBlock(id="photo", slot_id="image", source="upload", url="/aippt/assets/00000000-0000-0000-0000-000000000001", alt="实际配图", credit="原署名", locked=True)])
    deck = Deck(id="deck", title="图片导出", theme_id="ivory", slides=[slide])
    task = ClaimedTask(taskId="fixture", attempt=1, leaseToken="lease", fencingToken=1, type="export", catalogVersion=catalog_version(),
        payload={"font_manifest_hash":"unavailable", "snapshot":{"project":{"id":"deck","title":"图片导出","theme_id":"ivory"},"slides":[slide.model_dump(mode="json")]}})
    return deck, task


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["PNG", "JPEG", "WEBP"])
async def test_supported_private_images_export_as_real_native_pictures(fmt):
    output = BytesIO()
    image = Image.new("RGBA" if fmt == "WEBP" else "RGB", (300,200), (0,100,200,127) if fmt == "WEBP" else "blue")
    image.save(output, format=fmt, **({"icc_profile":b"fixture-profile"} if fmt == "WEBP" else {}))
    data = output.getvalue()
    if fmt == "WEBP": assert data[12:16] == b"VP8X"
    digest = hashlib.sha256(data).hexdigest()
    deck, task = deck_and_task(); center = Center(data)
    result, manifests = await TaskHandlers(center).run(task)
    assert result["verify"]["passed"] and result["font_mode"] == "estimate" and len(manifests) == 1
    presentation = Presentation(BytesIO(center.uploaded))
    pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1 and pictures[0].name == "aippt-image:photo"
    assert pictures[0].image.content_type == ("image/png" if fmt == "WEBP" else "image/"+fmt.lower())
    assert hashlib.sha256(center.data).hexdigest() == digest and deck.slides[0].blocks[-1].credit == "原署名"
    # 回读独立拒绝真实图片被静默占位或漏掉。
    presentation.slides[0].shapes._spTree.remove(pictures[0]._element)
    missing = BytesIO(); presentation.save(missing)
    report = verify_pptx(missing.getvalue(), deck)
    assert not report.passed and any(issue.check == "native_image" for issue in report.issues)


@pytest.mark.asyncio
async def test_corrupt_webp_never_becomes_successful_placeholder_export():
    deck, task = deck_and_task(); center = Center(b"RIFF"+b"\0"*4+b"WEBP"+b"bad"*20)
    with pytest.raises(ValueError, match="质量错误"): await TaskHandlers(center).run(task)
    assert center.uploaded is None


@pytest.mark.asyncio
@pytest.mark.parametrize("bad", [b"\x89PNG\r\n\x1a\n"+b"\0"*64, b"\xff\xd8\xff"+b"bad"*20, b"RIFF"+b"\0"*4+b"WEBP"+b"bad"*20])
async def test_bad_private_image_has_blocking_quality_issue_at_actual_page_and_slot(bad):
    _, task = deck_and_task(); task.type = "quality"; center = Center(bad)
    result, assets = await TaskHandlers(center).run(task)
    issues = [issue for issue in result["issues"] if issue["severity"] == "error" and "图片字节" in issue["message"]]
    assert not result["can_export"] and not assets and center.uploaded is None
    assert len(issues) == 1 and issues[0]["slide_id"] == "slide" and issues[0]["slot_id"] == "image"


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", ["lease", "unauthorized", "transport"])
async def test_quality_does_not_swallow_lease_permission_or_channel_failures(failure):
    import httpx
    from app.worker.protocol import LeaseLost
    error = LeaseLost("fixture") if failure == "lease" else httpx.ConnectError("fixture") if failure == "transport" else httpx.HTTPStatusError("fixture",request=httpx.Request("GET","http://center/"),response=httpx.Response(403))
    class FailedCenter(Center):
        async def read_asset(self, *args): raise error
    _, task = deck_and_task(); task.type = "quality"; center = FailedCenter(b"")
    with pytest.raises(type(error)): await TaskHandlers(center).run(task)
    assert center.uploaded is None
