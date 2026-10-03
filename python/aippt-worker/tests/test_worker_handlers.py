import copy
from io import BytesIO
import uuid
import pytest
from PIL import Image
from pptx import Presentation
from app.domain.sample import load_sample_deck
from app.domain.outline import OutlineDraft, OutlinePageDraft
from app.domain.slide_draft import SlideDraft, TextContent, BulletsContent
from app.domain.slide_patch import TextPatch, BulletsPatch
from app.worker.handlers import TaskHandlers
from app.worker.protocol import ClaimedTask
from app.worker.catalog import catalog_version


def task(kind, payload):
    return ClaimedTask(taskId="fixture", attempt=1, leaseToken="lease", fencingToken=1, type=kind, payload=payload,
                       catalogVersion=catalog_version())


class Center:
    def __init__(self): self.uploads=[]; self.reads=[]; self.assets={}
    async def upload_asset(self, task, data, filename, mime, role):
        self.uploads.append((data, filename, mime, role)); return {"asset_id": "image-1", "url": "/aippt/assets/image-1"}
    async def read_asset(self, task, asset_id, limit):
        self.reads.append(asset_id); return self.assets[asset_id]


class OutlineGenerator:
    async def generate(self, payload):
        return OutlineDraft(pages=[OutlinePageDraft(title=f"第{i}页", objective="明确变化原因", key_points=["事实一", "事实二"],
                       source_refs=[payload.sections[0].ref], layout_id="bullets") for i in range(payload.page_count)])


class SlideGenerator:
    async def generate(self, payload):
        return SlideDraft(blocks=[TextContent(slot_id="title", text=payload.page_title), BulletsContent(slot_id="body",
                    items=["交付耗时主要发生在需求与验收重复沟通，需统一正式记录", "统一中心任务版本防止迟到结果覆盖用户手工修改", "通过严格资源归属检查防止其他项目图片被读取"])])


class EditGenerator:
    async def generate(self, payload):
        return [TextPatch(block_id="t1", text="企图覆盖锁定标题"), BulletsPatch(block_id="b1", items=["合法修改一", "合法修改二"])]


class RelayoutGenerator:
    async def propose(self, **kwargs): return []


@pytest.mark.asyncio
async def test_relayout_uses_distinct_presets_without_content_changes():
    from app.domain.flex_edit import build_flex_tree_from_fixed
    from app.domain.layout import get_layout
    from app.domain.flex_layout import iter_leaf_block_ids
    deck=load_sample_deck(); slide=deck.slides[3]
    tree=build_flex_tree_from_fixed(get_layout(slide.layout_id),[b.model_dump(mode="json") for b in slide.blocks])
    slide=slide.model_copy(update={"layout_mode":"flex","layout_tree":tree})
    handlers=TaskHandlers(Center());handlers.generators["relayout"]=RelayoutGenerator()
    raw=slide.model_dump(mode="json"); original=copy.deepcopy(raw)
    output,_=await handlers.run(task("relayout",{"project":{"title":"样例"},"slide":raw}))
    assert raw==original and 1<=len(output["candidates"])<=3
    from app.domain.flex_layout import FlexContainer
    def no_override(node):
        return node.rect_override is None if node.type == "block" else all(no_override(child) for child in node.children)
    for candidate in output["candidates"]:
        parsed = FlexContainer.model_validate(candidate["layout_tree"])
        assert set(iter_leaf_block_ids(parsed)) == {b.id for b in slide.blocks}
        assert no_override(parsed)


@pytest.mark.asyncio
async def test_parse_text_stable_sections_repeat_without_database():
    handlers=TaskHandlers(Center()); payload={"source": {"id": "stable-src", "kind": "text", "text": "第一段\n\n第二段"}}
    output, assets = await handlers.run(task("parse_source", payload))
    again, _ = await handlers.run(task("parse_source", payload))
    assert {k:v for k,v in output.items() if k!="execution"} == {k:v for k,v in again.items() if k!="execution"} and len(output["sections"]) == 2 and not assets
    assert output["execution"]["usage_status"]=="not_applicable" and output["execution"]["elapsed_ms"]>=0
    assert all(s["id"] == s["ref"] and len(s["id"])==36 for s in output["sections"])


@pytest.mark.asyncio
async def test_source_document_private_asset_reads_and_process_parse():
    center=Center(); center.assets["doc"]="章节内容\n\n详细材料".encode()
    output, _ = await TaskHandlers(center).run(task("parse_source", {"source": {"id": "source-doc", "kind": "document", "filename": "材料.txt", "asset_id": "doc"}}))
    assert center.reads == ["doc"] and output["char_count"] > 0


@pytest.mark.asyncio
async def test_outline_and_slide_tasks_preserve_source_uuid_and_geometry():
    handlers=TaskHandlers(Center()); handlers.generators={"outline": OutlineGenerator(), "slide": SlideGenerator()}
    section_id=str(uuid.uuid4()); project={"id": "project", "title": "迁移计划", "tone": "professional", "page_count": 5, "layout_mode": "fixed", "theme_id": "ivory"}
    sources=[{"id": "source", "sections": [{"id": section_id, "text": "需统一中心与任务事实", "level": 0, "locator": "第一段"}]}]
    outline, _ = await handlers.run(task("generate_outline", {"project": project, "sources": sources}))
    assert len(outline["pages"]) == 5 and outline["pages"][0]["source_refs"] == [section_id]
    slide_id=str(uuid.uuid4())
    output, assets = await handlers.run(task("generate_slide", {"project": project, "sources": sources, "outline_page": outline["pages"][0], "slide": {"id": slide_id, "position": 1}}))
    assert len(output["slide"]["blocks"]) == 2 and len(output["geometry"]) == 2 and not assets
    assert output["slide"]["id"] == slide_id


@pytest.mark.asyncio
@pytest.mark.parametrize("history", [None, []])
async def test_ai_candidate_does_not_mutate_input_and_keeps_locked_title(history):
    center=Center(); handlers=TaskHandlers(center); handlers.generators["edit"]=EditGenerator()
    slide={"id":"slide", "revision":3, "title":"原标题", "layout_id":"bullets", "layout_mode":"fixed", "blocks":[
        {"id":"t1","slot_id":"title","type":"text","text":"原标题","locked":True},
        {"id":"b1","slot_id":"body","type":"bullets","items":["原要点一","原要点二"]}]}
    original=copy.deepcopy(slide)
    output,_=await handlers.run(task("ai_edit",{"project":{"title":"项目","tone":"professional"},"slide":slide,"instruction":"改短","history":history}))
    assert slide == original and output["candidate_slide"]["blocks"][0]["text"]=="原标题"
    assert output["discarded"] and output["operations"][0]["block_id"]=="b1"


@pytest.mark.asyncio
async def test_quality_and_native_export_snapshot_have_identical_source():
    deck=load_sample_deck(); snapshot={"project":{"id":deck.id,"title":deck.title,"theme_id":deck.theme_id},"slides":[s.model_dump(mode="json") for s in deck.slides]}
    original=copy.deepcopy(snapshot); center=Center(); handlers=TaskHandlers(center)
    quality,_=await handlers.run(task("quality",{"snapshot":snapshot}))
    assert quality["can_export"] and not center.uploads
    exported,assets=await handlers.run(task("export",{"snapshot":snapshot}))
    assert exported["verify"]["passed"] and snapshot==original
    assert assets[0]["role"]=="export" and assets[0]["size"]>1000
    presentation=Presentation(BytesIO(center.uploads[0][0]))
    assert len(presentation.slides)==len(deck.slides)
    assert any(shape.has_table for slide in presentation.slides for shape in slide.shapes)


@pytest.mark.asyncio
async def test_quality_missing_or_external_image_is_not_silently_exported():
    handlers=TaskHandlers(Center()); snapshot={"project":{"id":"p","title":"坏图","theme_id":"ivory"},"slides":[{
        "id":"s","layout_id":"image-left","blocks":[{"id":"image","slot_id":"image","type":"image","source":"stock","alt":"图","url":"https://private.invalid/image"}]}]}
    with pytest.raises(ValueError,match="私有图片"): await handlers.run(task("export",{"snapshot":snapshot}))
    assert not handlers.center.uploads


def test_image_pipeline_defaults_off_even_with_provider_credentials(monkeypatch):
    from app.core.config import Settings
    from app.images import pipeline
    cfg=Settings(_env_file=None,image_api_key="fixture-key",unsplash_access_key="fixture-key")
    assert cfg.image_enabled is False
    monkeypatch.setattr(pipeline,"get_settings",lambda:cfg)
    assert pipeline.create_image_pipeline(object()).enabled is False


@pytest.mark.asyncio
@pytest.mark.parametrize("local_enabled,permission",[(False,True),(True,None),(True,False),(True,1),(True,"true"),(False,False)])
async def test_image_resolution_requires_both_explicit_gates(monkeypatch,local_enabled,permission):
    from app.core.config import get_settings
    from app.domain.content import ImageBlock
    from app.worker import handlers as module
    cfg=get_settings();monkeypatch.setattr(cfg,"image_enabled",local_enabled)
    monkeypatch.setattr(cfg,"image_api_key","fixture-key");monkeypatch.setattr(cfg,"unsplash_access_key","fixture-key")
    def no_client(): raise AssertionError("未获双重许可时不应创建外部HTTP客户端")
    monkeypatch.setattr(module,"SafeAsyncClient",no_client)
    slide=next(s for s in load_sample_deck().slides if any(isinstance(b,ImageBlock) for b in s.blocks))
    center=Center();handlers=TaskHandlers(center)
    output,assets=await handlers._resolve_images(task("generate_slide",{"image_enabled":permission}),slide,"主题","页面")
    assert output is slide and assets==[] and center.uploads==[]


@pytest.mark.asyncio
async def test_image_resolution_enabled_gates_use_validated_fixture_and_asset_manifest(monkeypatch):
    from app.core.config import get_settings
    from app.domain.content import ImageBlock
    from app.images import pipeline
    from app.images.base import ImageAsset
    from app.worker import handlers as module
    monkeypatch.setattr(get_settings(),"image_enabled",True)
    data=BytesIO();Image.new("RGB",(20,10),"blue").save(data,format="PNG")
    calls=[]
    class Pipeline:
        async def fetch(self,request):
            calls.append(request);return ImageAsset(data=data.getvalue(),content_type="image/png",source="generated")
    class Client:
        async def __aenter__(self):return self
        async def __aexit__(self,*args):return None
    monkeypatch.setattr(module,"SafeAsyncClient",Client)
    monkeypatch.setattr(pipeline,"create_image_pipeline",lambda client:Pipeline())
    slide=next(s for s in load_sample_deck().slides if any(isinstance(b,ImageBlock) for b in s.blocks))
    center=Center()
    output,assets=await TaskHandlers(center)._resolve_images(task("generate_slide",{"image_enabled":True}),slide,"主题","页面")
    assert calls and len(center.uploads)==len(assets)==len(calls)
    assert all(a["role"]=="image" and a["size"]==len(data.getvalue()) for a in assets)
    assert any(isinstance(b,ImageBlock) and b.url=="/aippt/assets/image-1" and b.source=="generated" for b in output.blocks)


@pytest.mark.asyncio
async def test_image_gate_off_keeps_uploaded_asset_quality_and_native_export(monkeypatch):
    from app.core.config import get_settings
    from app.domain.content import ImageBlock
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    monkeypatch.setattr(get_settings(),"image_enabled",False)
    slide=next(s for s in load_sample_deck().slides if any(isinstance(b,ImageBlock) for b in s.blocks))
    blocks=[b.model_copy(update={"url":"/aippt/assets/uploaded-fixture","source":"upload","locked":True}) if isinstance(b,ImageBlock) else b for b in slide.blocks]
    slide=slide.model_copy(update={"blocks":blocks})
    data=BytesIO();Image.new("RGB",(800,450),"blue").save(data,format="PNG")
    center=Center();center.assets["uploaded-fixture"]=data.getvalue();handlers=TaskHandlers(center)
    snapshot={"project":{"id":"p","title":"已有上传图","theme_id":"ivory"},"slides":[slide.model_dump(mode="json")]}
    quality,_=await handlers.run(task("quality",{"snapshot":snapshot,"image_enabled":False}))
    exported,assets=await handlers.run(task("export",{"snapshot":snapshot,"image_enabled":False}))
    assert quality["can_export"] and exported["verify"]["passed"]
    assert center.reads==["uploaded-fixture","uploaded-fixture"]
    assert len(center.uploads)==1 and assets[0]["role"]=="export"
    presentation=Presentation(BytesIO(center.uploads[0][0]))
    assert any(shape.shape_type==MSO_SHAPE_TYPE.PICTURE for shape in presentation.slides[0].shapes)
