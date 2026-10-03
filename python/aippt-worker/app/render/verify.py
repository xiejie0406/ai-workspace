"""PPTX 回读与可编辑性验证。

输入渲染结果字节流与源 Deck，检查页数、画布、原生对象与内容完整性。
供导出接口与回归脚本复用，不依赖 HTTP 层。
"""

from __future__ import annotations

from io import BytesIO
from decimal import Decimal
import posixpath
import re
from typing import BinaryIO
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.presentation import Presentation as PresentationType
from pptx.shapes.base import BaseShape
from pptx.slide import Slide as PptxSlide
from pptx.util import Emu, Pt
from pydantic import BaseModel, Field

from app.domain.ambient import AMBIENT_SHAPE_PREFIX
from app.domain.content import (
    Block,
    BulletsBlock,
    CalloutBlock,
    CardsBlock,
    ChartBlock,
    Deck,
    ImageBlock,
    KpiBlock,
    Slide,
    TableBlock,
    TextBlock,
)
from app.domain.geometry import CANVAS_HEIGHT_PT, CANVAS_WIDTH_PT

# 半 pt 容差：渲染取整与自由形状边界可能有亚点误差
_BOUNDS_TOLERANCE_EMU = Emu(Pt(0.5))
# 覆盖整页的判据：宽高与面积均达到画布的 90% 以上
_FULL_PAGE_RATIO = 0.9
_CANVAS_TOLERANCE_EMU = Emu(Pt(0.5))
_CHART_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
             "x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
             "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


class _ChartDataInvalid(ValueError):
    """只携带固定诊断，不将数据正文或底层解析异常写进报告。"""


class VerifyIssue(BaseModel):
    """单条可编辑性/完整性问题，可定位到页与形状。"""

    check: str
    slide_index: int | None = None
    shape: str | None = None
    message: str


class VerifyReport(BaseModel):
    """回读验证报告。"""

    passed: bool
    issues: list[VerifyIssue] = Field(default_factory=list)
    slide_count: int = 0
    expected_slide_count: int = 0


def verify_pptx(data: bytes | BinaryIO, deck: Deck) -> VerifyReport:
    """用 python-pptx 重新打开导出结果并核对源 Deck。"""
    stream = BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
    if hasattr(stream, "seek"):
        stream.seek(0)

    issues: list[VerifyIssue] = []
    try:
        presentation = Presentation(stream)
    except Exception as error:
        return VerifyReport(
            passed=False,
            issues=[
                VerifyIssue(
                    check="open",
                    message=f"无法用 python-pptx 打开导出文件：{error}",
                )
            ],
            expected_slide_count=len(deck.slides),
        )

    expected = len(deck.slides)
    actual = len(presentation.slides)
    issues.extend(_check_page_count(actual, expected))
    issues.extend(_check_canvas_size(presentation))

    # 页数不一致时仍尽量检查已有页，避免一次失败掩盖其它问题
    for index, (source, pptx_slide) in enumerate(
        zip(deck.slides, presentation.slides, strict=False),
        start=1,
    ):
        issues.extend(_check_slide(index, source, pptx_slide, presentation))

    return VerifyReport(
        passed=not issues,
        issues=issues,
        slide_count=actual,
        expected_slide_count=expected,
    )


def _check_page_count(actual: int, expected: int) -> list[VerifyIssue]:
    if actual == expected:
        return []
    return [
        VerifyIssue(
            check="page_count",
            message=f"页数不一致：导出 {actual} 页，项目应为 {expected} 页",
        )
    ]


def _check_canvas_size(presentation: PresentationType) -> list[VerifyIssue]:
    expected_w = Emu(Pt(CANVAS_WIDTH_PT))
    expected_h = Emu(Pt(CANVAS_HEIGHT_PT))
    width = presentation.slide_width
    height = presentation.slide_height
    if width is None or height is None:
        return [
            VerifyIssue(
                check="canvas_size",
                message="页面尺寸缺失，无法确认是否为 16:9",
            )
        ]

    if abs(int(width) - int(expected_w)) > int(_CANVAS_TOLERANCE_EMU) or abs(
        int(height) - int(expected_h)
    ) > int(_CANVAS_TOLERANCE_EMU):
        ratio = (int(width) / int(height)) if int(height) else 0.0
        return [
            VerifyIssue(
                check="canvas_size",
                message=(
                    f"页面尺寸不是约定的 16:9（960×540 pt）："
                    f"当前约 {int(width) / 12700:.1f}×{int(height) / 12700:.1f} pt，"
                    f"宽高比 {ratio:.3f}"
                ),
            )
        ]
    return []


def _check_slide(
    slide_index: int,
    source: Slide,
    pptx_slide: PptxSlide,
    presentation: PresentationType,
) -> list[VerifyIssue]:
    issues: list[VerifyIssue] = []
    shapes = list(pptx_slide.shapes)

    issues.extend(_check_full_page_pictures(slide_index, shapes, presentation))
    issues.extend(_check_shape_bounds(slide_index, shapes, presentation))
    issues.extend(_check_native_tables(slide_index, source, shapes))
    issues.extend(_check_native_charts(slide_index, source, shapes))
    issues.extend(_check_native_images(slide_index, source, shapes))
    issues.extend(_check_text_in_frames(slide_index, source, shapes))
    issues.extend(_check_content_integrity(slide_index, source, shapes))
    issues.extend(_check_speaker_notes(slide_index, source, pptx_slide))
    return issues


def _check_speaker_notes(slide_index: int, source: Slide, pptx_slide: PptxSlide) -> list[VerifyIssue]:
    expected = source.speaker_notes or ""
    frame = pptx_slide.notes_slide.notes_text_frame if pptx_slide.has_notes_slide else None
    actual = frame.text if frame is not None else ""
    if actual == expected:
        return []
    return [VerifyIssue(check="speaker_notes", slide_index=slide_index, message=f"第 {slide_index} 页演讲备注与源快照不一致")]


def _check_full_page_pictures(
    slide_index: int,
    shapes: list[BaseShape],
    presentation: PresentationType,
) -> list[VerifyIssue]:
    slide_w = int(presentation.slide_width or 0)
    slide_h = int(presentation.slide_height or 0)
    if slide_w <= 0 or slide_h <= 0:
        return []

    slide_area = slide_w * slide_h
    issues: list[VerifyIssue] = []
    for shape in shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        width = int(shape.width)
        height = int(shape.height)
        area = width * height
        if (
            width >= slide_w * _FULL_PAGE_RATIO
            and height >= slide_h * _FULL_PAGE_RATIO
            and area >= slide_area * _FULL_PAGE_RATIO
        ):
            issues.append(
                VerifyIssue(
                    check="full_page_picture",
                    slide_index=slide_index,
                    shape=_shape_label(shape),
                    message=(
                        f"第 {slide_index} 页存在覆盖整页的图片"
                        f"（{_shape_label(shape)}），疑似用截图承载正文"
                    ),
                )
            )
    return issues


def _check_shape_bounds(
    slide_index: int,
    shapes: list[BaseShape],
    presentation: PresentationType,
) -> list[VerifyIssue]:
    slide_w = int(presentation.slide_width or 0)
    slide_h = int(presentation.slide_height or 0)
    tol = int(_BOUNDS_TOLERANCE_EMU)
    issues: list[VerifyIssue] = []

    for shape in shapes:
        # 氛围层是有意出血的装饰：光晕要靠画布边缘裁才不会被压变形
        if (shape.name or "").startswith(AMBIENT_SHAPE_PREFIX):
            continue
        left = int(shape.left)
        top = int(shape.top)
        right = left + int(shape.width)
        bottom = top + int(shape.height)
        if left < -tol or top < -tol or right > slide_w + tol or bottom > slide_h + tol:
            issues.append(
                VerifyIssue(
                    check="bounds",
                    slide_index=slide_index,
                    shape=_shape_label(shape),
                    message=(
                        f"第 {slide_index} 页形状越界：{_shape_label(shape)}"
                        f"（left={left}, top={top}, right={right}, bottom={bottom}；"
                        f"画布 {slide_w}×{slide_h} EMU）"
                    ),
                )
            )
    return issues


def _check_native_tables(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    expected = sum(1 for block in source.blocks if isinstance(block, TableBlock))
    if expected == 0:
        return []

    actual = sum(1 for shape in shapes if shape.has_table)
    if actual >= expected:
        return []
    return [
        VerifyIssue(
            check="native_table",
            slide_index=slide_index,
            message=(
                f"第 {slide_index} 页缺少原生表格：源内容有 {expected} 个表格块，"
                f"导出仅找到 {actual} 个 GraphicFrame 表格"
            ),
        )
    ]


def _check_native_charts(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    blocks = [block for block in source.blocks if isinstance(block, ChartBlock)]
    charts = [shape for shape in shapes if shape.has_chart]
    issues = []
    if len(charts) != len(blocks):
        issues.append(VerifyIssue(check="native_chart", slide_index=slide_index,
                      message=f"第 {slide_index} 页原生图表数量不一致：源 {len(blocks)} 个、导出 {len(charts)} 个"))
    # 原构造 fixed/flex 均按 source.blocks 顺序追加图表；逐一消费，不以相同文字复用形状。
    seen = set()
    for block, shape in zip(blocks, charts):
        try:
            chart = shape.chart
            identity = str(chart.part.partname)
            if identity in seen:
                raise _ChartDataInvalid("多个图表形状复用了同一图表数据部件")
            seen.add(identity)
            _verify_chart_data(chart, block)
        except Exception as error:
            reason = str(error) if isinstance(error, _ChartDataInvalid) else "图表或内嵌工作簿缺失、损坏或无法解析"
            issues.append(VerifyIssue(check="native_chart", slide_index=slide_index, shape=_shape_label(shape),
                          message=f"第 {slide_index} 页图表块 {block.id} 完整性失败：{reason}"))
    return issues


def _chart_read_budget(block: ChartBlock) -> int:
    """按已持有的源字符串及数据维度限制展开读取，允许合法长类别和最大维度。"""
    text_bytes = sum(len(value.encode("utf-8")) for value in block.categories)
    text_bytes += sum(len(series.name.encode("utf-8")) for series in block.series)
    cells = len(block.categories) * (len(block.series) + 1) + len(block.series)
    # 缓存重复类别；32倍涵盖 XML 转义/UTF-8，结构开销覆盖缓存及 XLSX 全部原包成员。
    return 2 * 1024 * 1024 + 32 * text_bytes * (len(block.series) + 1) + 1024 * cells


def _chart_xml(data: bytes, budget: int):
    if len(data) > budget or b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise _ChartDataInvalid("图表数据展开超出源内容预算或含非法 XML 声明")
    return ET.fromstring(data)


def _cache_values(ser, path: str, count: int) -> list[str]:
    cache = ser.find(path, _CHART_NS)
    if cache is None or cache.find("c:ptCount", _CHART_NS) is None:
        raise _ChartDataInvalid("可见图表缓存缺失")
    points = cache.findall("c:pt", _CHART_NS)
    if int(cache.find("c:ptCount", _CHART_NS).get("val")) != count or len(points) != count:
        raise _ChartDataInvalid("可见图表缓存维度不同")
    indexed = {}
    for point in points:
        index = int(point.get("idx"))
        value = point.find("c:v", _CHART_NS)
        if index in indexed or value is None:
            raise _ChartDataInvalid("可见图表缓存索引或值缺失")
        indexed[index] = value.text or ""
    if set(indexed) != set(range(count)):
        raise _ChartDataInvalid("可见图表缓存索引不同")
    return [indexed[index] for index in range(count)]


def _equal_chart_number(value: str, expected: float, *, workbook=False) -> bool:
    # python-pptx cache 用 float 字符串；锁定 XlsxWriter 写数值用 .16G。核既有序列化，不加误差容忍。
    actual = Decimal(value)
    source = Decimal(format(expected, ".16G") if workbook else str(expected))
    return actual.is_finite() and source.is_finite() and actual == source


def _column(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _formula_ref(value: str | None):
    match = re.fullmatch(r"('(?:[^']|'')+'|[^!]+)!\$?([A-Z]+)\$?([1-9]\d*)(?::\$?([A-Z]+)\$?([1-9]\d*))?", value or "")
    if match is None:
        raise _ChartDataInvalid("图表数据引用范围无效")
    sheet, col, row, end_col, end_row = match.groups()
    if sheet.startswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, col + row, (end_col + end_row if end_col else col + row)


def _verify_chart_data(chart, block: ChartBlock) -> None:
    budget = _chart_read_budget(block)
    root = _chart_xml(chart.part.blob, budget)
    plot = root.find("c:chart/c:plotArea", _CHART_NS)
    nodes = [node for node in plot if node.tag.rsplit("}", 1)[-1].endswith("Chart")]
    kind = {"bar": "barChart", "column": "barChart", "line": "lineChart", "pie": "pieChart"}[block.chart_type]
    if len(nodes) != 1 or nodes[0].tag != "{" + _CHART_NS["c"] + "}" + kind:
        raise _ChartDataInvalid("原生图表类型与源不同")
    node = nodes[0]
    if kind == "barChart":
        direction = node.find("c:barDir", _CHART_NS)
        if direction is None or direction.get("val") != ("bar" if block.chart_type == "bar" else "col"):
            raise _ChartDataInvalid("原生图表方向与源不同")
    visible = block.series[:1] if block.chart_type == "pie" else block.series
    series_nodes = node.findall("c:ser", _CHART_NS)
    if len(series_nodes) != len(visible):
        raise _ChartDataInvalid("可见图表系列数量不同")
    sheet_name, cells = _chart_workbook_cells(chart, root, budget)
    for index, (ser, expected) in enumerate(zip(series_nodes, visible)):
        for field in ("idx", "order"):
            position = ser.find("c:" + field, _CHART_NS)
            if position is None or int(position.get("val")) != index:
                raise _ChartDataInvalid("可见图表系列次序不同")
        if _cache_values(ser, "c:tx/c:strRef/c:strCache", 1) != [expected.name]:
            raise _ChartDataInvalid("可见图表系列名称不同")
        if _cache_values(ser, "c:cat/c:strRef/c:strCache", len(block.categories)) != block.categories:
            raise _ChartDataInvalid("可见图表类别不同")
        values = _cache_values(ser, "c:val/c:numRef/c:numCache", len(expected.values))
        if len(expected.values) != len(block.categories) or not all(_equal_chart_number(v, e) for v, e in zip(values, expected.values)):
            raise _ChartDataInvalid("可见图表数值或维度不同")
        column, last = _column(index + 2), str(len(block.categories) + 1)
        refs = [("c:tx/c:strRef/c:f", column + "1", column + "1"),
                ("c:cat/c:strRef/c:f", "A2", "A" + last),
                ("c:val/c:numRef/c:f", column + "2", column + last)]
        for path, start, end in refs:
            if _formula_ref(ser.findtext(path, namespaces=_CHART_NS)) != (sheet_name, start, end):
                raise _ChartDataInvalid("可见图表引用了错误的工作表或数据范围")
    expected_cells = {"A" + str(row): ("text", value) for row, value in enumerate(block.categories, 2) if value != ""}
    for index, series in enumerate(block.series):
        column = _column(index + 2)
        if series.name != "":
            expected_cells[column + "1"] = ("text", series.name)
        if len(series.values) != len(block.categories):
            raise _ChartDataInvalid("源图表数据维度不同")
        for row, value in enumerate(series.values, 2):
            expected_cells[column + str(row)] = ("number", value)
    if set(cells) != set(expected_cells):
        raise _ChartDataInvalid("内嵌工作簿完整数据单元格集合不同")
    for address, (kind, expected) in expected_cells.items():
        actual_kind, actual = cells[address]
        if actual_kind != kind or (actual != expected if kind == "text" else not _equal_chart_number(actual, expected, workbook=True)):
            raise _ChartDataInvalid("内嵌工作簿类别、系列名称或数值不同")


def _chart_workbook_cells(chart, chart_root, budget: int):
    external = chart_root.find("c:externalData", _CHART_NS)
    if external is None:
        raise _ChartDataInvalid("图表缺少内嵌工作簿关系")
    rid = external.get("{" + _CHART_NS["r"] + "}id")
    relation = chart.part.rels[rid]
    if relation.is_external or relation.reltype != "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package":
        raise _ChartDataInvalid("图表工作簿不是自身的内部嵌入关系")
    data = relation.target_part.blob
    if len(data) > budget:
        raise _ChartDataInvalid("内嵌工作簿超出源内容读取预算")
    with ZipFile(BytesIO(data)) as archive:
        members = archive.infolist()
        names = [item.filename for item in members]
        if len(members) > 64 or len(set(names)) != len(names) or sum(item.file_size for item in members) > budget:
            raise _ChartDataInvalid("内嵌工作簿成员重复或展开超过源内容预算")
        if any(item.flag_bits & 1 or item.filename.startswith("/") or ".." in item.filename.split("/") for item in members):
            raise _ChartDataInvalid("内嵌工作簿包含非法成员")
        def read_xml(path):
            info = archive.getinfo(path)
            if info.file_size > budget:
                raise _ChartDataInvalid("内嵌工作簿 XML 超过源内容预算")
            with archive.open(info) as stream:
                raw = stream.read(budget + 1)
            return _chart_xml(raw, budget)
        book = read_xml("xl/workbook.xml")
        rels = read_xml("xl/_rels/workbook.xml.rels")
        targets = {}
        for item in rels:
            identity, target = item.get("Id"), item.get("Target", "")
            if identity in targets or item.get("TargetMode") == "External" or "\\" in target:
                raise _ChartDataInvalid("工作簿关系重复、外链或非法")
            normalized = posixpath.normpath(posixpath.join("xl", target)).lstrip("/")
            if normalized.startswith("../") or not normalized.startswith("xl/"):
                raise _ChartDataInvalid("工作簿关系越界")
            targets[identity] = (item.get("Type"), normalized)
        sheets = book.findall("x:sheets/x:sheet", _CHART_NS)
        if len(sheets) != 1:
            raise _ChartDataInvalid("内嵌工作簿数据工作表数量不同")
        sheet = sheets[0]
        reltype, target = targets[sheet.get("{" + _CHART_NS["r"] + "}id")]
        if not reltype.endswith("/worksheet"):
            raise _ChartDataInvalid("内嵌工作表关系类型不同")
        strings = []
        shared = [path for typ, path in targets.values() if typ.endswith("/sharedStrings")]
        if len(shared) > 1:
            raise _ChartDataInvalid("共享字符串关系重复")
        if shared:
            strings = [_office_string("".join(si.itertext())) for si in read_xml(shared[0]).findall("x:si", _CHART_NS)]
        cells, seen = {}, set()
        for cell in read_xml(target).findall("x:sheetData/x:row/x:c", _CHART_NS):
            address = cell.get("r")
            if address in seen or not re.fullmatch(r"[A-Z]+[1-9]\d*", address or "") or cell.find("x:f", _CHART_NS) is not None:
                raise _ChartDataInvalid("内嵌工作簿单元格重复、地址无效或含公式")
            seen.add(address)
            kind = cell.get("t")
            value = cell.findtext("x:v", namespaces=_CHART_NS)
            if kind == "s":
                number = int(value)
                if not 0 <= number < len(strings):
                    raise _ChartDataInvalid("共享字符串索引无效")
                value, kind = strings[number], "text"
            elif kind == "inlineStr":
                inline = cell.find("x:is", _CHART_NS)
                value, kind = _office_string("".join(inline.itertext())), "text"
            elif kind == "str":
                value, kind = _office_string(value or ""), "text"
            elif kind in {None, "n"}:
                kind = "number"
            else:
                raise _ChartDataInvalid("内嵌工作簿数据单元格类型不同")
            # 原 writer 合法空字符串可能不写 cell；空格与真正缺失的非空内容仍严格区分。
            if value not in {None, ""}:
                cells[address] = (kind, value)
        return sheet.get("name"), cells


def _office_string(value: str) -> str:
    decoded = re.sub(r"_x([0-9a-fA-F]{4})_", lambda match: chr(int(match.group(1), 16)), value)
    return decoded.encode("utf-16-le", errors="surrogatepass").decode("utf-16-le", errors="surrogatepass")


def _check_native_images(slide_index: int, source: Slide, shapes: list[BaseShape]) -> list[VerifyIssue]:
    expected = [block for block in source.blocks if isinstance(block, ImageBlock) and block.url]
    if not expected:
        return []
    pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    names = {shape.name for shape in pictures}
    missing = [block.id for block in expected if f"aippt-image:{block.id}" not in names]
    if len(pictures) >= len(expected) and not missing:
        return []
    return [VerifyIssue(check="native_image", slide_index=slide_index,
                        message=f"第 {slide_index} 页真实图片缺失：源内容 {len(expected)} 张、导出 {len(pictures)} 张，缺失块 {','.join(missing)}")]


def _check_text_in_frames(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    """正文必须落在真正的文本框里，而不是位图。"""
    expected_texts = [
        text
        for block in source.blocks
        if isinstance(block, (TextBlock, BulletsBlock, KpiBlock))
        for text in key_texts_from_block(block)
    ]
    if not expected_texts:
        return []

    frame_texts = _collect_text_frame_strings(shapes)
    if not frame_texts:
        return [
            VerifyIssue(
                check="text_frame",
                slide_index=slide_index,
                message=(
                    f"第 {slide_index} 页有文字内容，但未找到带非空文字的文本框（可能被位图替代）"
                ),
            )
        ]

    issues: list[VerifyIssue] = []
    for text in expected_texts:
        if not _text_present(text, frame_texts):
            issues.append(
                VerifyIssue(
                    check="text_frame",
                    slide_index=slide_index,
                    message=(f"第 {slide_index} 页文字未出现在文本框中：「{_preview(text)}」"),
                )
            )
    return issues


def _check_content_integrity(
    slide_index: int,
    source: Slide,
    shapes: list[BaseShape],
) -> list[VerifyIssue]:
    haystack = _collect_all_strings(shapes)
    issues: list[VerifyIssue] = []
    for block in source.blocks:
        texts = key_texts_from_block(block)
        if isinstance(block, ChartBlock) and block.chart_type == "pie":
            # 饼图按原构造只绘首系列；全部系列必须另经本图工作簿完整性检查，不能混入可见 haystack。
            texts = [value for value in block.categories if value.strip()]
            texts.extend(series.name for series in block.series[:1] if series.name.strip())
        for text in texts:
            if _text_present(text, haystack):
                continue
            issues.append(
                VerifyIssue(
                    check="content_integrity",
                    slide_index=slide_index,
                    message=(
                        f"第 {slide_index} 页内容缺失：块 {block.id}"
                        f"（{block.type}）的关键文字「{_preview(text)}」未在导出结果中找到"
                    ),
                )
            )
    return issues


def key_texts_from_block(block: Block) -> list[str]:
    """块级关键文字，供完整性核对与回归脚本复用。"""
    match block:
        case TextBlock(text=text):
            return [text] if text.strip() else []
        case BulletsBlock(items=items):
            return [item for item in items if item.strip()]
        case KpiBlock(value=value, label=label, note=note):
            texts = [value, label]
            if note:
                texts.append(note)
            return [item for item in texts if item.strip()]
        case TableBlock(header=header, rows=rows):
            texts = list(header)
            for row in rows:
                texts.extend(row)
            return [item for item in texts if item.strip()]
        case ChartBlock(categories=categories, series=series):
            texts = list(categories)
            texts.extend(item.name for item in series if item.name)
            return [item for item in texts if item.strip()]
        case ImageBlock():
            return []
        case CardsBlock(items=items):
            texts: list[str] = []
            for item in items:
                if item.title.strip():
                    texts.append(item.title)
                if item.desc.strip():
                    texts.append(item.desc)
            return texts
        case CalloutBlock(text=text):
            return [text] if text.strip() else []
        case _:
            return []


def _collect_text_frame_strings(shapes: list[BaseShape]) -> list[str]:
    texts: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            continue
        if shape.has_table or shape.has_chart:
            continue
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            value = paragraph.text.strip()
            if value:
                texts.append(value)
    return texts


def _collect_all_strings(shapes: list[BaseShape]) -> list[str]:
    texts = list(_collect_text_frame_strings(shapes))
    for shape in shapes:
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value:
                        texts.append(value)
        if shape.has_chart:
            texts.extend(_chart_strings(shape))
    return texts


def _chart_strings(shape: BaseShape) -> list[str]:
    chart = shape.chart
    texts: list[str] = []
    try:
        for plot in chart.plots:
            for category in plot.categories:
                value = str(category).strip()
                if value:
                    texts.append(value)
    except Exception:
        pass
    try:
        for series in chart.series:
            name = str(series.name).strip() if series.name is not None else ""
            if name:
                texts.append(name)
    except Exception:
        pass
    return texts


def _text_present(expected: str, haystack: list[str]) -> bool:
    needle = expected.strip()
    if not needle:
        return True
    if needle in haystack:
        return True
    # 允许段落合并或单元格拼接后的包含关系
    return any(needle in item for item in haystack)


def _shape_label(shape: BaseShape) -> str:
    name = getattr(shape, "name", None) or "未命名形状"
    return str(name)


def _preview(text: str, limit: int = 40) -> str:
    compact = " ".join(text.split())
    if len(compact) <= limit:
        return compact
    return compact[: limit - 1] + "…"
